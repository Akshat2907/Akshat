"""
Quick visual QA renderer for PPTX files.
Draws boxes + text labels for each shape so we can check layout/overflow.
"""
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageDraw, ImageFont
import textwrap, sys, os

SLIDE_W_IN = 10.0
SLIDE_H_IN = 7.5
DPI = 120
PX_W = int(SLIDE_W_IN * DPI)
PX_H = int(SLIDE_H_IN * DPI)

def emu_to_px(emu, dpi=DPI):
    return int(emu / 914400 * dpi)

pptx_path = sys.argv[1] if len(sys.argv) > 1 else 'HAL_Strip_Profile.pptx'
prs = Presentation(pptx_path)

slide = prs.slides[0]

img = Image.new('RGB', (PX_W, PX_H), 'white')
draw = ImageDraw.Draw(img)

try:
    font_sm = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 11)
    font_xs = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 9)
    font_hd = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', 13)
except:
    font_sm = ImageFont.load_default()
    font_xs = font_sm
    font_hd = font_sm

for shape in slide.shapes:
    x = emu_to_px(shape.left)
    y = emu_to_px(shape.top)
    w = emu_to_px(shape.width)
    h = emu_to_px(shape.height)

    # Draw shape outline
    draw.rectangle([x, y, x+w, y+h], outline='#AAAAAA', width=1)

    # Extract text if present
    if shape.has_text_frame:
        full_text = ' '.join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
        if full_text:
            # Wrap text to fit width (approx 7px per char at font_sm)
            chars_per_line = max(10, w // 7)
            lines = textwrap.wrap(full_text, chars_per_line)[:6]
            txt = '\n'.join(lines)
            try:
                draw.text((x+3, y+2), txt, fill='#222222', font=font_xs)
            except Exception:
                pass

    # For tables, draw a grid hint
    if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
        table = shape.table
        draw.rectangle([x, y, x+w, y+h], outline='#CC2133', width=2)
        draw.text((x+3, y+2), f'TABLE {table.rows.__len__()}r x {table.columns.__len__()}c', fill='#CC2133', font=font_hd)

    # For charts
    if shape.shape_type == MSO_SHAPE_TYPE.CHART:
        draw.rectangle([x, y, x+w, y+h], outline='#0066CC', width=2)
        draw.text((x+3, y+2), '[CHART]', fill='#0066CC', font=font_hd)

# Draw slide boundary
draw.rectangle([0, 0, PX_W-1, PX_H-1], outline='black', width=2)

# Grid lines for quadrant reference
draw.line([(int(PX_W*0.51), 0), (int(PX_W*0.51), PX_H)], fill='#DDDDDD', width=1)
draw.line([(0, int(PX_H*0.495)), (PX_W, int(PX_H*0.495))], fill='#DDDDDD', width=1)

out = pptx_path.replace('.pptx', '_preview.png')
img.save(out)
print(f'Preview saved: {out}  [{PX_W}x{PX_H}px]')
